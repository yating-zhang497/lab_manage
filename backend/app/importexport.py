"""批量导入与导出：CSV/Excel/bib 解析、唯一标识upsert、Excel/BibTeX生成。"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal


def cell_to_str(v) -> str:
    """Excel 单元格 → 文本，避免数字/日期被 str() 破坏：
    - 整数值的 float 去掉 .0（学号 20240001.0 → 20240001，防科学计数法）
    - datetime/date → YYYY-MM-DD，否则 "2025-06-01 00:00:00" 无法被日期解析
    """
    if v is None:
        return ""
    if isinstance(v, bool):
        return "是" if v else ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        v = float(v)
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return format(v, "f").rstrip("0").rstrip(".")
    return str(v).strip()


# Excel 1900 历元（含 1900 伪闰年修正，序列号 >= 61 精确）
_EPOCH = date(1899, 12, 30)
# 只识别 5 位及以上数字（≈1927 年起），避开 4 位年份与经费等数值字段；
# 仅在已知日期字段上调用，数字字段（经费/年份）走各自的解析路径，不受影响
_EXCEL_SERIAL_MIN = 10000
_EXCEL_SERIAL_MAX = 80000


def excel_serial_to_iso(v):
    """Excel 日期序列号(数字或纯数字字符串) → 'YYYY-MM-DD'；非序列号返回 None。

    用户在 Excel 里输入 '2026年1月1日' / '2026/1/1' 等，Excel 会存为序列号(如 46023)，
    若单元格格式未被 openpyxl 识别为日期，cell_to_str 会得到 '46023' 文本；
    本函数仅在日期字段上将其还原为 ISO 日期，避免经费/年份等数值被误转。
    """
    try:
        n = int(v)
    except (ValueError, TypeError):
        return None
    if not (_EXCEL_SERIAL_MIN <= n <= _EXCEL_SERIAL_MAX):
        return None
    return (_EPOCH + timedelta(days=n)).isoformat()


def read_table(data: bytes, filename: str) -> list:
    """CSV(utf-8/gbk)或xlsx → list[dict]，表头自动去空白。"""
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = None
        for enc in ("utf-8-sig", "gb18030"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("无法识别CSV文件编码")
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        for row in reader:
            rows.append(
                {
                    _norm_header(k): (v or "").strip()
                    for k, v in row.items()
                    if k is not None
                }
            )
        return rows

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in all_rows[0]]
    rows = []
    for r in all_rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        rows.append(
            {
                _norm_header(header[i]): ("" if i >= len(r) else cell_to_str(r[i]))
                for i in range(len(header))
                if header[i]
            }
        )
    return rows


def _norm_header(s: str) -> str:
    return re.sub(r"\s+", "", str(s)).lower()


PERSON_FIELDS = {
    "uid": ["学号/工号", "学号", "工号", "学生学号", "uid"],
    "name": ["姓名", "学生姓名", "name"],
    "english_name": ["英文名", "英文姓名", "englishname"],
    "degree_type": ["类型", "学位类型", "学生类别"],
    "status": ["状态", "在读状态"],
    "current_activity": ["目前状态", "当前去向", "目前去向"],
    "research_direction": ["学位论文题目或研究方向", "研究方向", "学位论文题目"],
    "phone": ["电话号码", "电话", "手机号"],
    "email": ["email", "邮箱", "电子邮件"],
    "notes": ["备注"],
}

PROJECT_FIELDS = {
    "project_no": ["编号/索引项", "项目编号", "编号", "索引项"],
    "name": ["课题名称", "项目名称", "名称"],
    "category": ["类别", "项目类别"],
    "source": ["项目来源", "来源"],
    "status": ["状态"],
    "duration": ["项目起止时间", "项目起止", "起止时间", "起止"],
    "budget_no": ["经费编号"],
    "targets": ["指标/成果", "指标要求", "项目KPI", "项目kpi"],
    "targets_done": ["指标完成情况", "完成情况", "KPI完成情况", "kpi完成情况"],
    "budget_total": ["项目总经费（万元）", "项目总经费(万元)", "项目总经费", "总经费"],
    "budget_ours": ["我方总经费(万元)", "我方总经费（万元）", "我方总经费", "我方经费(万元)", "我方经费（万元）", "我方经费"],
    "budget_received": ["已到款经费(万元)", "已到款经费（万元）", "已到款经费", "到账经费(万元)", "到账经费（万元）", "到账经费"],
    "notes": ["备注", "其他说明"],
}

PAPER_FIELDS = {
    "title": ["论文题目", "标题", "title"],
    "authors_raw": ["所有作者名排序", "作者", "作者列表", "authors"],
    "venue": ["期刊/会议名称", "期刊/会议", "刊物", "venue", "journal"],
    "year": ["年份", "发表时间", "year"],
    "volume": ["卷号", "卷", "volume"],
    "issue": ["期号", "期", "issue", "number"],
    "pages": ["页码", "pages"],
    "ccf_class": ["ccf类别", "ccf"],
    "th_cpl_class": ["th-cpl类别"],
    "ack_index": ["收录检索号", "检索号", "论文检索号"],
    "acknowledgements": ["致谢", "acknowledgements"],
    "project_no_1": ["项目编号1", "基金标注1"],
    "project_no_2": ["项目编号2", "基金标注2"],
    "project_no_3": ["项目编号3", "基金标注3"],
    "project_seq_1": ["项目编号1对应排序", "项目编号1排序"],
    "project_seq_2": ["项目编号2对应排序", "项目编号2排序"],
    "project_seq_3": ["项目编号3对应排序", "项目编号3排序"],
    "evaluation": ["论文成果评价", "成果评价"],
    "notes": ["备注"],
    "url": ["论文链接", "链接", "url"],
    "demo_url": ["demo链接"],
    "doi": ["doi"],
}

IP_FIELDS = {
    "title": ["产权名称", "名称", "title"],
    "ip_type": ["类型"],
    "grant_date": ["获批时间", "授权日", "登记时间"],
    "applicant": ["专利权人", "申请人", "产权人"],
    "notes": ["其他说明", "备注"],
}

AWARD_FIELDS = {
    "award_name": ["奖项名称", "获奖名称"],
    "level": ["奖励等级"],
    "award_date": ["获奖时间"],
    "issuer": ["颁奖单位", "颁发单位"],
    "cert_no": ["证书编号"],
    "pr_url": ["公示链接"],
    "all_awardees": ["所有获奖人"],
    "summary": ["获奖作品"],
}

FLOAT_FIELDS = {"budget_total", "budget_ours", "budget_received"}
INT_FIELDS = {"year"}


def pick(row: dict, field: str, alias_map: dict):
    for alias in alias_map.get(field, []):
        v = row.get(alias) or row.get(alias.lower())
        if v not in (None, ""):
            return v
    for key, v in row.items():
        if key.lower() in [a.lower() for a in alias_map.get(field, [])] and v not in (None, ""):
            return v
    return None


def row_to_entity(row: dict, alias_map: dict, extra_fields: dict = None, extra_date_keys: set = None) -> dict:
    """将一行 Excel 数据转换为实体字段字典。
    extra_fields: 自定义字段 key → 别名列表，会把匹配的值收集到 out["extra"]。
    extra_date_keys: 自定义字段中 dtype=date 的 key 集合，其值会做 Excel 序列号→ISO 日期还原。
    """
    out = {}
    extra = {}
    for field in alias_map:
        v = pick(row, field, alias_map)
        if v is None:
            continue
        if field in FLOAT_FIELDS:
            try:
                out[field] = float(str(v).replace(",", ""))
            except ValueError:
                continue
        elif field in INT_FIELDS:
            try:
                out[field] = int(float(str(v)))
            except ValueError:
                continue
        else:
            out[field] = str(v)
    # 自定义字段（不在内置 alias_map 中，通过 extra_fields 指定别名）
    if extra_fields:
        for field, aliases in extra_fields.items():
            v = pick(row, field, {field: aliases})
            if v is None:
                continue
            if extra_date_keys and field in extra_date_keys:
                iso = excel_serial_to_iso(v)
                extra[field] = iso if iso else str(v)
            else:
                extra[field] = str(v)
    if extra:
        out["extra"] = extra
    return out


def split_authors(authors_raw: str) -> list:
    """"Last, First; Last, First" → 有序作者名列表。"""
    if not authors_raw:
        return []
    names = [n.strip() for n in re.split(r"[;；]", authors_raw) if n.strip()]
    return names


def build_paper_bib(p) -> str:
    if p.bib_raw:
        return p.bib_raw.strip()
    first_author = ""
    if p.authors_raw:
        parts = split_authors(p.authors_raw)
        if parts:
            first_author = parts[0].split(",")[0].strip()
    word = re.sub(r"[^A-Za-z]", "", (p.title or "")[:20])[:8] or "paper"
    key = p.bib_key or f"{first_author.lower() or 'anon'}{p.year or ''}{word.lower()}"
    lines = [f"@{'inproceedings' if p.venue_type == '会议' else 'article'}{{{key},"]
    fields = [
        ("title", p.title),
        ("author", (p.authors_raw or "").replace("; ", " and ")),
        ("journal" if p.venue_type != "会议" else "booktitle", p.venue),
        ("year", p.year),
        ("volume", p.volume),
        ("number", p.issue),
        ("pages", p.pages),
        ("address", p.address),
        ("publisher", p.publisher),
        ("doi", p.doi),
        ("url", p.url),
        ("abstract", p.abstract),
        ("note", p.notes),
    ]
    for k, v in fields:
        if v not in (None, ""):
            lines.append(f"  {k} = {{{v}}},")
    lines.append("}")
    return "\n".join(lines)


def make_excel(headers: list, rows: list) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(
            12, min(50, len(str(h)) * 2 + 8)
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_template(cols: list) -> bytes:
    """生成带下拉选项与日期格式的 Excel 导入模板。
    cols: list[(key, label, options_or_None, dtype_or_None)]（兼容三元组）；
    select 字段自动加下拉验证，date 字段整列设 yyyy-mm-dd 格式，使 Excel 把
    用户输入的任意日期写法存为序列号+日期格式，openpyxl 读回 datetime → ISO，
    避免落入 '46023' 序列号文本。
    """
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation

    headers = [c[1] if len(c) > 1 else "" for c in cols]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)

    # select 字段加下拉验证；date 字段整列设日期格式
    for col_idx, col in enumerate(cols, start=1):
        options = col[2] if len(col) > 2 else None
        dtype = col[3] if len(col) > 3 else None
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        if options:
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(options) + '"',
                allow_blank=True,
                showDropDown=False,
            )
            dv.sqref = f"{col_letter}2:{col_letter}1048576"
            ws.add_data_validation(dv)
        elif dtype == "date":
            for r in range(2, 201):
                ws[f"{col_letter}{r}"].number_format = "yyyy-mm-dd"

    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(
            12, min(50, len(str(h)) * 2 + 8)
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
