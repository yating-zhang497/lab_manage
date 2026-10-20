from conftest import make_person

import io

FIELD_BODY = {"entity": "people", "label": "本科学校", "dtype": "text"}


def test_builtin_select_options_editable(client, auth):
    # 字段管理：修改内置下拉字段的选项应保存生效
    r = client.put(
        "/api/fields/builtin/people/status",
        json={"options": ["在读/在职", "毕业/出站", "其他"]},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    fields = client.get("/api/fields?entity=people", headers=auth).json()
    status = [f for f in fields if f["key"] == "status"][0]
    assert status["options"] == ["在读/在职", "毕业/出站", "其他"]

    # 改名 + 改选项同时提交
    r = client.put(
        "/api/fields/builtin/people/degree_type",
        json={"label": "培养类型", "options": ["硕士", "博士", "博士后", "科研助理", "其他"]},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    fields = client.get("/api/fields?entity=people", headers=auth).json()
    degree = [f for f in fields if f["key"] == "degree_type"][0]
    assert degree["label"] == "培养类型"
    assert degree["options"][-1] == "其他"


def test_template_follows_field_defs(client, auth):
    import openpyxl

    # 模板按当前字段定义即时生成：改名后模板表头同步变化
    client.put(
        "/api/fields/builtin/projects/name",
        json={"label": "课题全称"},
        headers=auth,
    )
    r = client.get("/api/import/template/projects", headers=auth)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    headers = [c.value for c in next(wb.active.iter_rows(max_row=1))]
    assert "课题全称" in headers
    assert "项目编号" in headers

    r = client.get("/api/import/template/papers", headers=auth)
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    headers = [c.value for c in next(wb.active.iter_rows(max_row=1))]
    assert headers[0] == "论文题目"
    assert "作者1" in headers and "项目编号1" in headers


def test_orphan_field_defs_not_listed(client, auth):
    # 已下线内置字段的旧覆盖记录不应复活成自定义字段
    from conftest import TestSession
    from app.models import FieldDef

    db = TestSession()
    try:
        db.add(FieldDef(entity="projects", key="abbr", label="简称"))
        db.add(FieldDef(entity="ip", key="status", label="状态"))
        db.commit()
    finally:
        db.close()

    projects = client.get("/api/fields?entity=projects", headers=auth).json()
    assert "abbr" not in {f["key"] for f in projects}
    ip_fields = client.get("/api/fields?entity=ip", headers=auth).json()
    assert "status" not in {f["key"] for f in ip_fields}


def test_field_crud(client, auth):
    r = client.post("/api/fields", json=FIELD_BODY, headers=auth)
    assert r.status_code == 201, r.text
    f = r.json()
    assert f["key"].startswith("cf_")

    # 同名标签允许重复创建（key自动生成不同），但显式指定相同key会冲突
    r = client.post(
        "/api/fields", json={**FIELD_BODY, "key": f["key"]}, headers=auth
    )
    assert r.status_code == 409

    r = client.put(f"/api/fields/{f['id']}", json={"label": "本科院校"}, headers=auth)
    assert r.json()["label"] == "本科院校"

    listed = client.get("/api/fields?entity=people", headers=auth).json()
    assert any(x["id"] == f["id"] for x in listed)


def test_extra_value_roundtrip_and_export(client, auth):
    f = client.post("/api/fields", json=FIELD_BODY, headers=auth).json()
    select_field = client.post(
        "/api/fields",
        json={"entity": "people", "label": "培养方式", "dtype": "select", "options": ["统考", "推免"]},
        headers=auth,
    ).json()

    p = make_person(client, auth, uid="F001", name="字段测试")
    r = client.put(
        f"/api/people/{p['id']}",
        json={"extra": {f["key"]: "清华大学", select_field["key"]: "推免", "not_defined": "x"}},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["extra"][f["key"]] == "清华大学"
    assert r.json()["extra"][select_field["key"]] == "推免"
    assert "not_defined" not in r.json()["extra"]

    cols = client.get("/api/export/columns?type=people", headers=auth).json()
    keys = {c["key"] for c in cols}
    assert f"extra:{f['key']}" in keys

    prev = client.post(
        "/api/export/preview",
        json={"type": "people", "columns": ["name", f"extra:{f['key']}"], "filters": {}},
        headers=auth,
    ).json()
    row = [x for x in prev["rows"] if x["name"] == "字段测试"][0]
    assert row[f"extra:{f['key']}"] == "清华大学"

    r = client.put(f"/api/fields/{f['id']}", json={"enabled": False}, headers=auth)
    assert r.json()["enabled"] is False
    disabled_cols = client.get("/api/export/columns?type=people", headers=auth).json()
    assert f"extra:{f['key']}" not in {c["key"] for c in disabled_cols}

    r = client.delete(f"/api/fields/{f['id']}", headers=auth)
    assert r.json()["ok"] is True
