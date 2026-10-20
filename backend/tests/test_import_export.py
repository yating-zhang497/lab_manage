from conftest import make_paper, make_person, make_project

PEOPLE_CSV = """学号/工号,姓名,类型,状态,目前状态,Email
D001,陈梦璇,博士,在读/在职,,cmx@example.com
D002,张金潇,博士,毕业/出站,企业实习,jxz@example.com
,无学号人员,硕士,,,
"""

PROJECTS_CSV = """课题名称,项目类别,项目起止时间,状态,编号/索引项,项目来源,经费编号,项目总经费（万元）,我方总经费(万元),已到款经费(万元)
燕山石化VOCs溯源方法研究,横向,-2025.12,已结题,HS001,北京华清,JJ001,50.00,50,45
消费级显卡大模型推理服务通信优化技术研究,横向,-2026.1,在研,HS002,深信服科技股份有限公司,JJ002,50.00,50,25
"""

PAPERS_CSV = """论文题目,所有作者名排序,期刊/会议名称,年份,卷号,期号,页码,CCF类别
Deep Learning for Weather,"Fu, Haohuan; Zhang, San",Science,2024,10,2,1-20,A
,缺标题论文,Nature,2024,,,,
"""

BIB_SAMPLE = """
@article{fu2024weather,
  title = {AI Weather Prediction at Scale},
  author = {Fu, Haohuan and Zhang, San},
  journal = {Nature},
  year = {2024},
  volume = {12},
  pages = {1--15},
  doi = {10.1000/test}
}
@inproceedings{zhang2023hpc,
  title = {HPC Optimization Study},
  author = {Zhang, San},
  booktitle = {SC23},
  year = {2023}
}
"""


def csv_file(name, content):
    return {"file": (name, content.encode("utf-8"), "text/csv")}


def test_import_people_upsert(client, auth):
    r = client.post("/api/import/people", files=csv_file("p.csv", PEOPLE_CSV), headers=auth)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["ok_count"] == 2
    assert rep["fail_count"] == 1
    assert "学号/工号" in rep["errors"][0]

    people = client.get("/api/people", headers=auth).json()
    assert len(people) == 2

    updated = PEOPLE_CSV.replace("cmx@example.com", "new@example.com")
    r2 = client.post("/api/import/people", files=csv_file("p.csv", updated), headers=auth)
    assert r2.json()["ok_count"] == 2
    people = client.get("/api/people", headers=auth).json()
    assert len(people) == 2
    assert any(p["email"] == "new@example.com" for p in people)


def test_import_projects_duration_and_budget_no(client, auth):
    r = client.post("/api/import/projects", files=csv_file("p.csv", PROJECTS_CSV), headers=auth)
    assert r.status_code == 200, r.text
    assert r.json()["ok_count"] == 2
    projects = client.get("/api/projects", headers=auth).json()
    one = [p for p in projects if p["project_no"] == "HS001"][0]
    assert one["duration"] == "-2025.12"
    assert one["budget_no"] == "JJ001"
    assert one["budget_total"] == 50.0


def test_import_papers_matches_authors(client, auth):
    fu = make_person(client, auth, name="付昊桓", uid="T001", english_name="Fu, Haohuan")
    r = client.post("/api/import/papers", files=csv_file("p.csv", PAPERS_CSV), headers=auth)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["ok_count"] == 1
    assert rep["fail_count"] == 1

    papers = client.get("/api/papers", headers=auth).json()
    assert len(papers) == 1
    detail = client.get(f"/api/papers/{papers[0]['id']}", headers=auth).json()
    assert len(detail["authors"]) == 2
    assert detail["authors"][0]["person_id"] == fu["id"]
    assert detail["authors"][1]["person_id"] is None


def test_import_bib(client, auth):
    r = client.post(
        "/api/import/papers/bib",
        files={"file": ("x.bib", BIB_SAMPLE.encode("utf-8"), "application/x-bibtex")},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["ok_count"] == 2
    papers = client.get("/api/papers", headers=auth).json()
    titles = {p["title"] for p in papers}
    assert {"AI Weather Prediction at Scale", "HPC Optimization Study"} <= titles
    conf = [p for p in papers if p["title"] == "HPC Optimization Study"][0]
    detail = client.get(f"/api/papers/{conf['id']}", headers=auth).json()
    assert detail["venue_type"] == "会议"


BBT_SAMPLE = """
@inproceedings{fu2025better,
  title        = {Better BibTeX Compatibility Test},
  author       = {Fu, Haohuan and Ye, Zhiyu},
  booktitle    = {SC25},
  date         = {2025-11-17},
  pages        = {1--10},
  address      = {St. Louis, MO, USA},
  publisher    = {IEEE},
  doi          = {10.1109/SC25.00001},
  url          = {https://doi.org/10.1109/SC25.00001},
  abstract     = {A test of Better BibTeX/BibLaTeX field capture.},
  note         = {Best paper award}
}
"""


def test_import_bib_better_bibtex_fields(client, auth):
    """Better BibTeX/BibLaTeX 导出：date 取年份+出版日期、journaltitle、
    以及 GB/T 7714 字段（出版地/出版者/摘要/链接/备注）都能解析带回。"""
    r = client.post(
        "/api/import/papers/bib",
        files={"file": ("bbt.bib", BBT_SAMPLE.encode("utf-8"), "application/x-bibtex")},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["ok_count"] == 1
    papers = client.get("/api/papers", headers=auth).json()
    p = [x for x in papers if x["title"] == "Better BibTeX Compatibility Test"][0]
    detail = client.get(f"/api/papers/{p['id']}", headers=auth).json()
    # date 字段 → 年份 + 完整出版日期
    assert detail["year"] == 2025
    assert detail["pub_date"] == "2025-11-17"
    # GB/T 7714 字段
    assert detail["venue"] == "SC25"
    assert detail["venue_type"] == "会议"
    assert detail["address"] == "St. Louis, MO, USA"
    assert detail["publisher"] == "IEEE"
    assert detail["doi"] == "10.1109/SC25.00001"
    assert detail["url"] == "https://doi.org/10.1109/SC25.00001"
    assert "Better BibTeX" in detail["abstract"]
    assert detail["notes"] == "Best paper award"


def test_import_papers_author_columns_ordered(client, auth):
    fu = make_person(client, auth, name="付昊桓", uid="T001", english_name="Fu, Haohuan")
    csv = (
        "论文题目,作者1,作者2,作者3,期刊/会议名称,发表时间\n"
        'Author Column Paper,"Zhang, San","Fu, Haohuan",Li Si,Science,2025\n'
    )
    r = client.post("/api/import/papers", files=csv_file("p.csv", csv), headers=auth)
    assert r.status_code == 200, r.text
    assert r.json()["ok_count"] == 1
    paper = client.get("/api/papers", headers=auth).json()[0]
    detail = client.get(f"/api/papers/{paper['id']}", headers=auth).json()
    assert [a["author_name"] for a in detail["authors"]] == [
        "Zhang, San", "Fu, Haohuan", "Li Si",
    ]
    assert detail["authors"][0]["seq"] == 1
    assert detail["authors"][1]["person_id"] == fu["id"]

    # 作者1/作者2 导出列 + 项目编号对应排序列
    r = client.post(
        "/api/export/preview",
        json={"type": "papers", "columns": ["title", "author_1", "author_2"], "ids": [paper["id"]]},
        headers=auth,
    )
    row = r.json()["rows"][0]
    assert row["author_1"] == "Zhang, San"
    assert row["author_2"] == "Fu, Haohuan"


def test_import_awards(client, auth):
    make_person(client, auth, name="王五", uid="D100")
    csv = (
        "奖项名称,奖励等级,获奖时间,获奖作品,获奖人1\n"
        "学术奖学金,校级,2025-06-01,优秀论文,王五\n"
    )
    r = client.post("/api/import/awards", files=csv_file("a.csv", csv), headers=auth)
    assert r.status_code == 200, r.text
    assert r.json()["ok_count"] == 1
    awards = client.get("/api/awards", headers=auth).json()
    assert len(awards) == 1
    a = awards[0]
    assert a["award_name"] == "学术奖学金"
    assert a["level"] == "校级"
    assert a["summary"] == "优秀论文"
    assert len(a["recipients"]) == 1
    assert a["recipients"][0]["recipient_name"] == "王五"
    assert a["recipients"][0]["person_id"] is not None


def test_import_papers_matches_chinese_name(client, auth):
    # 作者用中文名录入时也能匹配人员库的"姓名"
    make_person(client, auth, name="王五", uid="T777", english_name="Wang, Wu")
    csv = "论文题目,作者1,作者2,期刊/会议名称,发表时间\n中文作者论文,王五,赵六,Science,2025\n"
    r = client.post("/api/import/papers", files=csv_file("p.csv", csv), headers=auth)
    assert r.status_code == 200, r.text
    paper = client.get("/api/papers", headers=auth).json()[0]
    detail = client.get(f"/api/papers/{paper['id']}", headers=auth).json()
    assert detail["authors"][0]["author_name"] == "王五"
    assert detail["authors"][0]["person_id"] is not None
    assert detail["authors"][1]["person_id"] is None


def test_precise_match_does_not_collide_on_abbrev(client, auth):
    """精确匹配：'Wang, Wei' 不得命中 'Wang, Wenqiang'（旧缩写变体 wang we 造成的误匹配）。"""
    from conftest import TestSession
    from app.routers.imports import _build_person_lookup, _match_person_in_lookup

    make_person(client, auth, name="王文强", uid="T001", english_name="Wang, Wenqiang")
    make_person(client, auth, name="王威", uid="T002", english_name="Wang, Wei")
    db = TestSession()
    lookup = _build_person_lookup(db)
    # "Wang, Wei" 应精确命中王威(T002)，而不是王文强(T001)
    p = _match_person_in_lookup(lookup, "Wang, Wei")
    assert p is not None and p.uid == "T002"
    # 缩写署名 "Wang WW" 不再自动命中任何人（需显式录入姓名变体）
    assert _match_person_in_lookup(lookup, "Wang WW") is None
    db.close()


def test_rematch_authors_full_clears_and_rematches(client, auth):
    """rematch-authors(full) 清空全部旧关联后用精确匹配重绑。"""
    from conftest import TestSession
    from app.models import PaperAuthor

    fu = make_person(client, auth, name="付昊桓", uid="T001", english_name="Fu, Haohuan")
    p = make_paper(client, auth, title="Rematch Paper",
                   authors=[{"author_name": "Fu, Haohuan", "seq": 1}])
    # create 不自动匹配，person_id 初始为空
    assert client.get(f"/api/papers/{p['id']}", headers=auth).json()["authors"][0]["person_id"] is None
    # 人为把关联置错（绑到不存在的 person），full 重匹配应纠正
    db = TestSession()
    db.query(PaperAuthor).filter(PaperAuthor.paper_id == p["id"]).update({PaperAuthor.person_id: 999999})
    db.commit(); db.close()
    r = client.post("/api/papers/rematch-authors?full=true", headers=auth)
    assert r.status_code == 200, r.text
    assert r.json()["matched"] >= 1
    assert r.json()["cleared"] >= 1
    detail = client.get(f"/api/papers/{p['id']}", headers=auth).json()
    assert detail["authors"][0]["person_id"] == fu["id"]


def test_import_ip_with_inventor_columns(client, auth):
    fu = make_person(client, auth, name="付昊桓", uid="T001", english_name="Fu, Haohuan")
    csv = (
        "产权名称,类型,获批时间,专利权人,发明人1,发明人2,其他说明\n"
        '一种数据处理方法,专利,2025-06-01,清华大学,"Fu, Haohuan",Zhang San,\n'
        "某系统V1.0,软著,,清华大学,,,\n"
    )
    r = client.post("/api/import/ip", files=csv_file("ip.csv", csv), headers=auth)
    assert r.status_code == 200, r.text
    assert r.json()["ok_count"] == 2

    ips = client.get("/api/ip", headers=auth).json()
    assert len(ips) == 2
    patent = [x for x in ips if x["title"] == "一种数据处理方法"][0]
    assert patent["ip_type"] == "patent"
    assert patent["grant_date"] == "2025-06-01"
    assert patent["inventors"][0]["person_id"] == fu["id"]
    assert patent["inventors"][0]["is_first"] is True
    soft = [x for x in ips if x["title"] == "某系统V1.0"][0]
    assert soft["ip_type"] == "software_copyright"

    # 重复导入：按产权名称upsert，不新增
    r2 = client.post("/api/import/ip", files=csv_file("ip.csv", csv), headers=auth)
    assert r2.json()["ok_count"] == 2
    assert len(client.get("/api/ip", headers=auth).json()) == 2


def test_excel_serial_date_recovery():
    """Excel 日期序列号应还原为 ISO 日期；4位年份/经费等数值不被误转。"""
    from app.importexport import excel_serial_to_iso

    assert excel_serial_to_iso(46023) == "2026-01-01"   # 2026-01-01 的 Excel 序列号
    assert excel_serial_to_iso("46023") == "2026-01-01"
    # 4 位年份不当作序列号（避免与“仅年份”输入冲突）
    assert excel_serial_to_iso(2026) is None
    assert excel_serial_to_iso("2026") is None
    # 超出范围 / 非数字
    assert excel_serial_to_iso(99999) is None
    assert excel_serial_to_iso("abc") is None


def test_import_ip_excel_serial_date(client, auth):
    """Excel 把日期存为序列号(General 格式)时，导入应还原为 ISO 日期而非 '46023'。"""
    files = xlsx_file(
        "ip.xlsx",
        ["产权名称", "类型", "获批时间"],
        [["一种数据处理方法", "专利", 46023]],   # 46023 = 2026-01-01 的 Excel 序列号
    )
    r = client.post("/api/import/ip", files=files, headers=auth)
    assert r.status_code == 200, r.text
    ips = client.get("/api/ip", headers=auth).json()
    assert ips[0]["grant_date"] == "2026-01-01"


def test_import_ip_requires_title_and_type(client, auth):
    csv = "产权名称,类型\n,专利\n无类型条目,未知类型\n"
    r = client.post("/api/import/ip", files=csv_file("ip.csv", csv), headers=auth)
    rep = r.json()
    assert rep["ok_count"] == 0
    assert rep["fail_count"] == 2


def test_import_templates_downloadable(client, auth):
    for url in (
        "/api/import/template/people",
        "/api/import/template/projects",
        "/api/import/template/papers",
        "/api/import/template/ip",
    ):
        r = client.get(url, headers=auth)
        assert r.status_code == 200
        assert r.content[:2] == b"PK"
    r = client.get("/api/import/template/papers/bib", headers=auth)
    assert r.status_code == 200
    assert b"@article" in r.content


def test_export_people_xlsx(client, auth):
    make_person(client, auth, uid="X001", name="测试甲")
    r = client.post(
        "/api/export/custom",
        json={"type": "people", "columns": [], "filters": {}},
        headers=auth,
    )
    assert r.status_code == 200
    assert r.content[:2] == b"PK"


def test_export_paper_builtin_fields_not_dropped(client, auth):
    """导出必须包含所选内置字段值，不能只导出列表可见字段。
    回归：_paper_row 曾硬编码字段列表，遗漏 corresponding_text/extract_status，
    导致即使数据库有值也导出为空（被误判为'只导出了列表可见字段'）。"""
    make_paper(
        client, auth, title="Regression Paper",
        corresponding_text="Zhang San", extract_status="已提取",
        volume="9", issue="1", pages="22",
    )
    r = client.post(
        "/api/export/preview",
        json={
            "type": "papers",
            "columns": ["title", "corresponding_text", "extract_status", "volume", "issue", "pages"],
            "filters": {"keyword": "Regression Paper"},
        },
        headers=auth,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total"] == 1
    row = body["rows"][0]
    # 所选字段全部有值（不只是列表可见的 title）
    assert row["title"] == "Regression Paper"
    assert row["corresponding_text"] == "Zhang San"
    assert row["extract_status"] == "已提取"
    assert row["volume"] == "9"
    assert row["issue"] == "1"
    assert row["pages"] == "22"


def test_custom_export_cross_table(client, auth):
    fu = make_person(client, auth, name="付昊桓", uid="T001", english_name="Fu, Haohuan")
    project = make_project(client, auth, project_no="T2125006", name="NSFC项目")
    r = client.post(
        "/api/papers",
        json={
            "title": "Custom Export Paper",
            "year": 2025,
            "authors": [
                {"author_name": "Fu, Haohuan", "person_id": fu["id"], "seq": 1, "is_corresponding": True}
            ],
            "funding": [{"project_id": project["id"], "funding_no": "T2125006", "seq": 1}],
        },
        headers=auth,
    )
    assert r.status_code == 201, r.text

    cols = client.get("/api/export/columns?type=papers", headers=auth).json()
    keys = {c["key"] for c in cols}
    assert {"title", "team_uids", "project_nos", "corresponding_names"} <= keys

    # 作者维度多值列可一起导出
    r = client.post(
        "/api/export/preview",
        json={
            "type": "papers",
            "columns": ["title", "team_uids", "team_names", "corresponding_names"],
            "filters": {"person_id": fu["id"]},
        },
        headers=auth,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total"] == 1
    row = body["rows"][0]
    assert row["team_uids"] == "T001"
    assert row["team_names"] == "付昊桓"
    assert row["corresponding_names"] == "付昊桓"

    # 项目维度多值列单独导出
    r = client.post(
        "/api/export/preview",
        json={
            "type": "papers",
            "columns": ["title", "project_nos"],
            "filters": {"person_id": fu["id"]},
        },
        headers=auth,
    )
    assert r.status_code == 200, r.text
    assert r.json()["rows"][0]["project_nos"] == "T2125006"

    # 同时选择两个维度的多值列 → 拒绝导出
    r = client.post(
        "/api/export/preview",
        json={
            "type": "papers",
            "columns": ["title", "team_names", "project_nos"],
            "filters": {"person_id": fu["id"]},
        },
        headers=auth,
    )
    assert r.status_code == 422
    assert "无法导出" in r.json()["detail"]

    r = client.post(
        "/api/export/custom",
        json={"type": "people", "columns": ["uid", "name", "paper_count"], "filters": {}},
        headers=auth,
    )
    assert r.status_code == 200
    assert r.content[:2] == b"PK"

    r = client.post(
        "/api/export/custom",
        json={"type": "papers", "columns": ["bad_column"], "filters": {}},
        headers=auth,
    )
    assert r.status_code == 422


def test_cross_table_export_papers(client, auth):
    fu = make_person(client, auth, name="付昊桓", uid="T001", english_name="Fu, Haohuan")
    stranger = make_person(client, auth, name="外部人员", uid="T002")
    project = make_project(client, auth, project_no="T2125006")
    r = client.post(
        "/api/papers",
        json={
            "title": "Cross Table Paper",
            "year": 2025,
            "venue": "SC",
            "authors": [{"author_name": "Fu, Haohuan", "person_id": fu["id"], "seq": 1}],
            "funding": [{"project_id": project["id"], "funding_no": "T2125006", "seq": 1}],
        },
        headers=auth,
    )
    assert r.status_code == 201, r.text
    client.post(
        "/api/papers",
        json={"title": "Unrelated Paper", "year": 2024},
        headers=auth,
    )

    r = client.post(
        "/api/export/custom",
        json={
            "type": "papers",
            "columns": ["title", "team_names"],
            "filters": {"person_id": fu["id"], "project_id": project["id"], "year_min": 2025, "year_max": 2026},
        },
        headers=auth,
    )
    assert r.status_code == 200
    assert r.content[:2] == b"PK"

    r2 = client.post(
        "/api/export/preview",
        json={"type": "papers", "columns": ["title"], "filters": {"person_id": stranger["id"]}},
        headers=auth,
    )
    assert r2.json()["total"] == 0

    r3 = client.get(f"/api/export/papers.bib?person_id={fu['id']}", headers=auth)
    text = r3.content.decode("utf-8")
    assert "Cross Table Paper" in text
    assert "Unrelated Paper" not in text


def xlsx_file(name, headers, rows):
    """构造真实 xlsx（保留单元格原始类型），用于验证数字/日期单元格的读取。"""
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return {
        "file": (
            name,
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


def test_import_people_xlsx_numeric_uid_and_date(client, auth):
    """Excel 里学号是数字、日期单元格时不应丢精度/丢日期。"""
    import datetime

    files = xlsx_file(
        "p.xlsx",
        ["学号/工号", "姓名"],
        [
            [20240001, "赵六"],
            [20240001.0, "赵六"],
        ],
    )
    r = client.post("/api/import/people", files=files, headers=auth)
    assert r.status_code == 200, r.text
    people = client.get("/api/people", headers=auth).json()
    assert [p["uid"] for p in people] == ["20240001"]


def test_import_people_custom_field(client, auth):
    """字段管理新增的自定义字段列应能随 Excel 导入并写入 extra。

    回归：row_to_entity 调用 pick() 时漏传 field 参数导致整请求 500。
    """
    import io

    import openpyxl

    # 1. 新增一个自定义字段
    f = client.post(
        "/api/fields",
        json={"entity": "people", "label": "本科学校", "dtype": "text"},
        headers=auth,
    ).json()
    assert f["key"].startswith("cf_")

    # 2. 下载真实模板（自定义字段列已包含在内）
    r = client.get("/api/import/template/people", headers=auth)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "本科学校" in headers

    # 3. 填一行数据，含自定义字段值
    rowvals = {}
    for h in headers:
        if h == "学号/工号":
            rowvals[h] = "CF001"
        elif h == "姓名":
            rowvals[h] = "自定义字段导入"
        elif h == "本科学校":
            rowvals[h] = "清华大学"
        else:
            rowvals[h] = None
    ws.append([rowvals.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)

    # 4. 导入：不应 500，且 ok_count==1
    files = {
        "file": (
            "p.xlsx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    r = client.post("/api/import/people", files=files, headers=auth)
    assert r.status_code == 200, r.text
    assert r.json()["ok_count"] == 1, r.text

    # 5. 自定义字段值应写入 extra 并可读回
    people = client.get("/api/people", headers=auth).json()
    p = [x for x in people if x["uid"] == "CF001"][0]
    detail = client.get(f"/api/people/{p['id']}", headers=auth).json()
    assert detail["extra"][f["key"]] == "清华大学"
